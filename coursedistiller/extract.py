"""Extract searchable text from downloaded attachments -> work/attachments_text/<stem>.md.
PDF (pdftotext + tesseract OCR fallback), xlsx, pptx, docx, csv. Requires poppler + tesseract on PATH."""
import glob
import os
import subprocess
import sys
import tempfile


def _sh(cmd, timeout=300):
    return subprocess.run(cmd, capture_output=True, timeout=timeout).stdout.decode("utf-8", "replace")


def _pdf(path, ocr=True):
    txt = _sh(["pdftotext", "-layout", path, "-"])
    try:
        info = subprocess.run(["pdfinfo", path], capture_output=True, timeout=30).stdout.decode("utf-8", "replace")
        pages = int([l for l in info.splitlines() if l.startswith("Pages:")][0].split(":")[1])
    except Exception:
        pages = max(1, txt.count("\f") + 1)
    if len(txt.strip()) >= 40 * pages or not ocr:
        return txt
    with tempfile.TemporaryDirectory() as td:
        try:
            subprocess.run(["pdftoppm", "-png", "-r", "200", path, os.path.join(td, "p")], capture_output=True, timeout=600)
            ocr_txt = "\n".join(_sh(["tesseract", img, "-", "--psm", "1"], timeout=180)
                                for img in sorted(glob.glob(os.path.join(td, "p*.png")))).strip()
            return ocr_txt if len(ocr_txt) > len(txt) else txt
        except Exception:
            return txt


def _xlsx(path):
    code = ("import openpyxl,sys\nwb=openpyxl.load_workbook(sys.argv[1],data_only=True,read_only=True)\n"
            "for ws in wb.worksheets:\n print('##',ws.title)\n for row in ws.iter_rows(values_only=True):\n"
            "  c=[('' if x is None else str(x)) for x in row]\n  if any(c): print('| '+' | '.join(c)+' |')\n")
    return subprocess.run([sys.executable, "-c", code, path], capture_output=True, timeout=180).stdout.decode("utf-8", "replace")


def _pptx(path):
    code = ("import sys\nfrom pptx import Presentation\np=Presentation(sys.argv[1])\n"
            "for i,s in enumerate(p.slides,1):\n print(f'\\n--- Slide {i} ---')\n"
            " for sh in s.shapes:\n  if sh.has_text_frame: print(sh.text_frame.text)\n")
    return subprocess.run([sys.executable, "-c", code, path], capture_output=True, timeout=180).stdout.decode("utf-8", "replace")


def _docx(path):
    try:
        return _sh(["textutil", "-convert", "txt", "-stdout", path])  # macOS
    except Exception:
        import docx
        return "\n".join(p.text for p in docx.Document(path).paragraphs)


def extract(cfg, progress=lambda **k: None):
    if not cfg.attachments["extract_text"]:
        return 0
    files = [f for f in glob.glob(os.path.join(cfg.attachments_dir, "*")) if os.path.isfile(f)]
    done = 0
    for i, f in enumerate(files):
        stem = os.path.splitext(os.path.basename(f))[0]
        out = os.path.join(cfg.attachments_text, stem + ".md")
        if os.path.exists(out) and os.path.getsize(out) > 0:
            done += 1
            continue
        ext = f.lower().rsplit(".", 1)[-1]
        try:
            txt = {"pdf": lambda: _pdf(f, cfg.attachments["ocr_fallback"]), "xlsx": lambda: _xlsx(f),
                   "pptx": lambda: _pptx(f), "docx": lambda: _docx(f),
                   "csv": lambda: open(f, encoding="utf-8", errors="replace").read(),
                   "txt": lambda: open(f, encoding="utf-8", errors="replace").read()}.get(ext, lambda: None)()
            if txt and len(txt.strip()) >= 5:
                open(out, "w").write(f"# {os.path.basename(f)}\n\n" + txt.strip())
                done += 1
        except Exception:
            pass
        progress(phase="extract", done=i + 1, total=len(files), msg=f"{done} extracted")
    return done
